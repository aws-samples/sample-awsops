import io
import re

import pytest

from diagnosis import exporters

_SAMPLE = "# AWS 진단 리포트\n\n> 생성 일시: 2026-06-17 09:00 (KST)\n\n## 요약\n\n본문 문단입니다.\n\n- 항목 A\n- 항목 B\n\n| 키 | 값 |\n|----|----|\n| a  | 1  |\n"


def _doc(md):
    from docx import Document

    return Document(io.BytesIO(exporters.to_docx(md)))


def test_to_docx_returns_docx_zip_bytes():
    out = exporters.to_docx(_SAMPLE)
    assert isinstance(out, (bytes, bytearray)) and len(out) > 0
    assert bytes(out[:4]) == b"PK\x03\x04"  # DOCX is a zip (OOXML)


def test_to_docx_preserves_content():
    import io
    from docx import Document

    doc = Document(io.BytesIO(exporters.to_docx(_SAMPLE)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "AWS 진단 리포트" in text   # heading rendered
    assert "본문 문단입니다." in text   # body paragraph rendered
    assert any(t.rows for t in doc.tables)  # the markdown table became a docx table


@pytest.fixture
def pdf_browser_available():
    api = pytest.importorskip("playwright.sync_api")
    try:
        with api.sync_playwright() as p:
            browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
            browser.close()
    except (api.Error, OSError, AttributeError) as e:
        pytest.skip(f"chromium unavailable: {e}")


def test_to_pdf_returns_pdf_bytes(pdf_browser_available):
    # Once the browser is available, renderer failures must fail the test, not skip.
    out = exporters.to_pdf(_SAMPLE)
    assert isinstance(out, (bytes, bytearray)) and bytes(out[:5]) == b"%PDF-"


@pytest.fixture
def pdf_resource_server():
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Event, Thread

    requests, connections = [], []
    received = Event()

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "2")
            self.end_headers()
            self.wfile.write(b"OK")
            received.set()

        def log_message(self, *args):
            pass

    class Server(ThreadingHTTPServer):
        def get_request(self):
            request, address = super().get_request()
            connections.append(address)  # Count speculative TCP connects even without HTTP.
            return request, address

    try:
        server = Server(("127.0.0.1", 0), Handler)
    except OSError as e:
        pytest.skip(f"loopback server unavailable: {e}")
    thread = Thread(target=lambda: server.serve_forever(poll_interval=0.05), daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", requests, connections, received
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


@pytest.mark.parametrize("markup", [
    "![image]({url}/markdown-image)",
    '<img src="{url}/image">',
    '<style>@import url("{url}/import.css"); body {{ background: url("{url}/background") }}</style>',
    '<iframe src="{url}/frame"></iframe>',
    '<iframe srcdoc=\'<img src="{url}/nested-image">\'></iframe>',
    pytest.param('<link rel="prefetch" href="{url}/prefetch">', id="prefetch"),
    pytest.param('<LiNk ReL="pre&#102;etch" HrEf="{url}/encoded">', id="encoded-prefetch"),
    pytest.param('<link rel="preload" as="image" href="{url}/preload">', id="preload"),
    pytest.param('<link rel="preconnect" href="{url}">', id="preconnect"),
    pytest.param('<link rel="dns-prefetch" href="{url}">', id="dns-prefetch"),
    pytest.param(
        '<meta http-equiv="Content-Security-Policy" content="default-src *">'
        '<link rel="prefetch" href="{url}/permissive-policy">',
        id="permissive-policy-prefetch",
    ),
    pytest.param(
        '<iframe srcdoc=\'<link rel="prefetch" href="{url}/nested-prefetch">\'></iframe>',
        id="nested-prefetch",
    ),
    pytest.param(
        '<script src="{url}/script.js"></script>'
        '<img src="data:invalid" onerror="fetch(\'{url}/onerror\')">',
        id="active-markup",
    ),
    pytest.param('<meta http-equiv="refresh" content="0;url={url}/refresh">', id="refresh"),
])
def test_to_pdf_prevents_real_resource_requests(pdf_browser_available, pdf_resource_server, markup):
    from playwright.sync_api import sync_playwright

    url, requests, connections, received = pdf_resource_server
    # Prefetch is a positive control for the actual browser-level bypass, not just ordinary images.
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
        try:
            browser.new_page(java_script_enabled=False).set_content(
                f'<link rel="prefetch" href="{url}/control">')
            assert received.wait(timeout=5), "unguarded prefetch did not reach the loopback server"
        finally:
            browser.close()
    assert requests == ["/control"] and connections
    requests.clear()
    connections.clear()

    out = exporters.to_pdf("# Local report\n\n" + markup.format(url=url))
    assert out.startswith(b"%PDF-")
    assert requests == [], f"PDF renderer fetched injected resources: {requests}"
    assert connections == [], f"PDF renderer opened speculative connections: {connections}"


@pytest.mark.parametrize("markup", [
    '<link rel="prefetch" href="{url}/prefetch">',
    '</body></html><head><link rel="prefetch" href="{url}/new-head"></head>',
    '<meta http-equiv="Content-Security-Policy" content="default-src *">'
    '<link rel="prefetch" href="{url}/permissive-policy">',
])
def test_html_policy_blocks_prefetch_without_offline_or_routing(
        pdf_browser_available, pdf_resource_server, markup):
    from playwright.sync_api import sync_playwright

    url, requests, connections, _received = pdf_resource_server
    # Bypass the markup filter to exercise CSP independently, with no offline mode or interception.
    html = exporters._html("# Local report").replace("</body>", markup.format(url=url) + "</body>", 1)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
        try:
            page = browser.new_page(java_script_enabled=False)
            page.set_content(html)
            assert page.pdf().startswith(b"%PDF-")
        finally:
            browser.close()
    assert requests == [] and connections == []


_ACTIVE_MARKUP = [
    '<LiNk ReL="preconnect" href="{url}">'
    '<link rel="dns-prefetch prefetch" href="{url}/hint">',
    '<meta http-equiv="x-dns-prefetch-control" content="on">'
    '<meta http-equiv="refresh" content="0;url={url}/refresh">'
    '<base href="{url}/"><link rel="prefetch" href="relative">',
    '<meta http-equiv="Content-Security-Policy" content="default-src *">'
    '<iframe srcdoc=\'<link rel="preconnect" href="{url}">\'></iframe>',
    '<svg><foreignObject><link rel="prefetch" href="{url}/foreign"></foreignObject></svg>'
    '<script>document.title="active"</script><object data="{url}/object"></object>',
    '<style>/* </style/foo><link rel="prefetch" href="{url}/raw-text"> */</style>',
    '<a href="java&#115;cript:alert(1)" ping="{url}/ping" onclick="alert(1)">link</a>'
    '<img src="{url}/image" srcset="{url}/srcset 2x" onerror="alert(1)">',
]


@pytest.mark.parametrize("markup", _ACTIVE_MARKUP)
def test_to_pdf_removes_active_elements_before_rendering(
        pdf_browser_available, pdf_resource_server, monkeypatch, markup):
    from playwright.sync_api import Page

    url, requests, connections, _received = pdf_resource_server
    original_pdf = Page.pdf

    def inspect_and_print(page, *args, **kwargs):
        assert page.locator("link, base, iframe, script, object, embed, svg, math").count() == 0
        assert page.locator("[onclick], [onerror], [srcdoc], [srcset], [ping]").count() == 0
        assert page.locator("a[href^='javascript:'], img[src^='http']").count() == 0
        # Only the trusted charset and CSP metadata may reach the browser.
        assert page.locator("meta").count() == 2
        assert page.locator("meta[http-equiv='Content-Security-Policy']").count() == 1
        return original_pdf(page, *args, **kwargs)

    monkeypatch.setattr(Page, "pdf", inspect_and_print)
    assert exporters.to_pdf("# Local report\n\n" + markup.format(url=url)).startswith(b"%PDF-")
    assert requests == [] and connections == []


def test_to_pdf_preserves_inline_report_rendering(pdf_browser_available, monkeypatch):
    from playwright.sync_api import Page

    original_pdf = Page.pdf
    inspected = []

    def inspect_and_print(page, *args, **kwargs):
        assert page.locator("h1").inner_text() == "AWS 진단 리포트"
        assert page.locator("table td").all_text_contents() == ["a", "1"]
        assert page.locator("h1").is_visible()
        assert page.locator("table").evaluate("el => getComputedStyle(el).borderCollapse") == "collapse"
        assert "Noto Sans CJK KR" in page.locator("body").evaluate("el => getComputedStyle(el).fontFamily")
        assert page.locator("#inline").evaluate("el => getComputedStyle(el).color") == "rgb(18, 52, 86)"
        assert page.locator("#inline").evaluate(
            "el => getComputedStyle(el, '::after').content") == '"A & B > C"'
        assert page.locator("#pixel").evaluate("el => el.complete && el.naturalWidth === 1")
        assert page.locator("pre code").inner_text().strip() == '<link rel="preconnect" href="#example">'
        assert page.locator("link").count() == 0  # A code example stays text, never a resource hint.
        assert page.locator("a").get_attribute("href") == "#section"
        inspected.append(True)
        return original_pdf(page, *args, **kwargs)

    # Inspect the real page at the PDF boundary, retaining its actual context and security controls.
    monkeypatch.setattr(Page, "pdf", inspect_and_print)
    out = exporters.to_pdf(_SAMPLE + (
        '\n<style>#inline { color: #123456 }'
        '#inline::after { content: "A & B > C" }</style><p id="inline">본문</p>'
        '<img id="pixel" src="data:image/gif;base64,'
        'R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBRAA7">'
        '\n\n```html\n<link rel="preconnect" href="#example">\n```\n\n[section](#section)'
    ))
    assert inspected and out.startswith(b"%PDF-")


def test_html_template_uses_system_font_no_external_import():
    html = exporters._html("# t\n\n본문")
    assert "Noto Sans CJK KR" in html
    assert "@import" not in html and "fonts.googleapis" not in html  # no egress in private subnet


def test_docx_page_setup_a4_explicit_margins():
    # Round to whole mm: OOXML stores page size in integer twips, so Mm(210) round-trips as
    # ~210.0086mm (the standard A4-in-twips rounding), not an exact EMU match.
    sec = _doc(_SAMPLE).sections[0]
    assert round(sec.page_width.mm) == 210
    assert round(sec.page_height.mm) == 297
    assert round(sec.top_margin.mm) == 18
    assert round(sec.bottom_margin.mm) == 18
    assert round(sec.left_margin.mm) == 16
    assert round(sec.right_margin.mm) == 16


def test_docx_normal_style_korean_font_and_east_asia():
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    st = _doc(_SAMPLE).styles["Normal"]
    assert st.font.name == "Malgun Gothic"
    assert st.font.size == Pt(10.5)
    assert st.font.color.rgb == RGBColor.from_string("1F1E1D")
    rfonts = st.element.rPr.find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:eastAsia")) == "Malgun Gothic"


def test_docx_heading_styles_not_word_default():
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    styles = _doc(_SAMPLE).styles
    expect = [
        ("Heading 1", Pt(20), "14130F"),
        ("Heading 2", Pt(15), "14130F"),
        ("Heading 3", Pt(12.5), "8E4830"),  # brand-deep — distinguishes LLM ### subheads
    ]
    for name, size, color in expect:
        st = styles[name]
        assert st.font.size == size, name
        assert st.font.color.rgb == RGBColor.from_string(color), name
        rfonts = st.element.rPr.find(qn("w:rFonts"))
        assert rfonts is not None and rfonts.get(qn("w:eastAsia")) == "Malgun Gothic", name


def test_docx_h1_has_brand_bottom_rule():
    from docx.oxml.ns import qn

    st = _doc(_SAMPLE).styles["Heading 1"]
    pbdr = st.element.pPr.find(qn("w:pBdr"))
    assert pbdr is not None
    bottom = pbdr.find(qn("w:bottom"))
    assert bottom is not None
    assert bottom.get(qn("w:val")) == "single"
    assert bottom.get(qn("w:color")) == "D97757"


def test_docx_page_break_before_second_h2_only():
    md = "# t\n\n## A\n\nbody\n\n## B\n\nbody\n"
    doc = _doc(md)
    h2s = [p for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert len(h2s) == 2
    assert not h2s[0].paragraph_format.page_break_before
    assert h2s[1].paragraph_format.page_break_before is True


def test_docx_table_explicit_widths_everywhere():
    from docx.shared import Mm

    t = _doc(_SAMPLE).tables[0]
    assert t.autofit is False
    # Round to whole twips-of-a-mm: OOXML stores widths in integer twips, so an exact EMU
    # equality fails by the same rounding as the page-size check above.
    w_mm = round((Mm(178) // len(t.columns)) / 36000)  # 1mm == 36000 EMU
    for col in t.columns:
        assert round(col.width.mm) == w_mm
    for row in t.rows:
        for cell in row.cells:
            assert round(cell.width.mm) == w_mm


def test_docx_table_header_bold_and_shaded():
    from docx.oxml.ns import qn

    t = _doc(_SAMPLE).tables[0]  # `| 키 | 값 |` header + separator + one data row
    header = t.rows[0]
    for cell in header.cells:
        assert cell.paragraphs[0].runs and all(r.bold for r in cell.paragraphs[0].runs)
        shd = cell._tc.tcPr.find(qn("w:shd"))
        assert shd is not None
        assert shd.get(qn("w:fill")) == "F5DCCF"
    data_row = t.rows[1]
    for cell in data_row.cells:
        shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
        assert shd is None


def test_docx_table_no_header_gets_no_bold_or_shading():
    from docx.oxml.ns import qn

    md = "| a | b |\n| c | d |\n"  # two data rows, no `|---|---|` separator → headerless
    t = _doc(md).tables[0]
    for row in t.rows:
        for cell in row.cells:
            assert not any(r.bold for r in cell.paragraphs[0].runs)
            shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
            assert shd is None


def test_docx_table_calm_border_color():
    from docx.oxml.ns import qn

    t = _doc(_SAMPLE).tables[0]
    assert t.style.name == "Table Grid"
    top = t._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:top"))
    assert top.get(qn("w:color")) == "D7D3C7"
    assert top.get(qn("w:sz")) == "4"


def test_docx_toc_label_real_bold_no_literal_asterisks():
    md = "# t\n\n**목차**\n\n- [항목](#a)\n"
    doc = _doc(md)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in full_text
    toc_label = next(p for p in doc.paragraphs if p.text == "목차")
    assert toc_label.runs and all(r.bold for r in toc_label.runs)


def test_docx_inline_code_monospace_run():
    md = "# t\n\n- `inventory`: ok\n"
    doc = _doc(md)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "`" not in full_text
    item = next(p for p in doc.paragraphs if p.text == "inventory: ok")
    code_run = next(r for r in item.runs if r.text == "inventory")
    assert code_run.font.name == "Consolas"
    plain_run = next(r for r in item.runs if r.text == ": ok")
    assert plain_run.font.name != "Consolas"


def test_docx_underscore_line_italic_no_literal_underscores():
    from docx.shared import RGBColor

    md = "# t\n\n_이 섹션 생성에 실패했습니다 (degraded): boom_\n"
    doc = _doc(md)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "_" not in full_text
    p = next(p for p in doc.paragraphs if "degraded" in p.text)
    assert p.runs and all(r.italic for r in p.runs)
    assert p.runs[0].font.color.rgb == RGBColor.from_string("5F5A4D")


def test_docx_blockquote_muted_color():
    from docx.shared import RGBColor

    p = next(p for p in _doc(_SAMPLE).paragraphs if "생성 일시" in p.text)
    assert p.runs[0].italic is True
    assert p.runs[0].font.color.rgb == RGBColor.from_string("5F5A4D")


def test_docx_ordered_list_uses_list_number_style():
    md = "# t\n\n1. 첫째\n2. 둘째\n"
    doc = _doc(md)
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(numbered) == 2
    for p in numbered:
        assert not re.match(r"^\d+[.)]", p.text)  # raw "1." must not leak into the text


def test_docx_fenced_code_block_rendered_not_literal():
    from docx.oxml.ns import qn

    md = "# t\n\n```\naws s3 ls\n```\n\n본문\n"
    doc = _doc(md)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "```" not in full_text
    code_p = next(p for p in doc.paragraphs if p.text == "aws s3 ls")
    assert code_p.runs[0].font.name == "Consolas"
    shd = code_p._p.pPr.find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == "F7F6F2"
    # the paragraph after the fence must still render normally (fence consumption doesn't
    # swallow trailing content)
    assert any(p.text == "본문" for p in doc.paragraphs)


def test_docx_unclosed_fence_does_not_crash_or_leak_backticks():
    md = "# t\n\n```\naws s3 ls\n"  # opens a fence, never closes it (truncation)
    doc = _doc(md)  # must not raise
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "```" not in full_text
    # must be real fence handling, not an accidental artifact of _add_runs' naive backtick check
    code_p = next(p for p in doc.paragraphs if p.text == "aws s3 ls")
    assert code_p.runs[0].font.name == "Consolas"


@pytest.mark.parametrize("markup", _ACTIVE_MARKUP)
def test_pdf_markup_filter_without_browser_removes_active_html(markup):
    # No pdf_browser_available fixture: CI exercises the pure filter without Chromium.
    from lxml import html

    tree = html.fromstring(exporters._html("# Local report\n\n" + markup.format(url="https://fixture.invalid")))
    assert not tree.xpath("//link|//base|//iframe|//script|//object|//embed|//svg|//math")
    assert not tree.xpath("//*[@onclick or @onerror or @srcdoc or @srcset or @ping]")
    assert not tree.xpath("//a[starts-with(@href, 'javascript:')]|//img[starts-with(@src, 'http')]")
    assert len(tree.xpath("//meta")) == 2
    assert tree.xpath("//meta[@http-equiv='Content-Security-Policy']/@content") == [exporters._PDF_CSP]
    assert "Local report" in tree.text_content()


def test_pdf_markup_filter_without_browser_preserves_formatting_and_escaped_examples():
    from lxml import html

    tree = html.fromstring(exporters._html(_SAMPLE + (
        '\n<style>#inline { color: #123456 } #inline::after { content: "A & B > C" }</style>'
        '<p id="inline" style="font-weight: bold">본문 &amp; text</p>'
        '<img id="pixel" src="data:image/gif;base64,R0lGODlhAQABAIAAAAAAAP///w==">'
        '\n\n```html\n<link rel="prefetch" href="#example">\n```\n\n[section](#section)'
    )))
    assert tree.xpath("//h1") and tree.xpath("//table")
    assert tree.xpath("//p[@id='inline']/@style") == ["font-weight: bold"]
    assert 'content: "A & B > C"' in tree.xpath("//style")[-1].text
    assert tree.xpath("//img[@id='pixel']/@src")[0].startswith("data:image/gif;base64,")
    assert '<link rel="prefetch" href="#example">' in tree.xpath("//pre/code")[0].text
    assert not tree.xpath("//link")
    assert tree.xpath("//a/@href") == ["#section"]


def test_pdf_markup_unmatched_closers_preserve_allowed_tree_without_browser():
    # Deep allowed nesting plus unmatched allowed closers used to re-scan the stack each time.
    body = "<div>" * 2000 + "</span>" * 2000 + "kept</div>" + "</div>" * 1999
    assert exporters._PdfMarkup().render(body) == "<div>" * 2000 + "kept" + "</div>" * 2000


def test_pdf_markup_cdata_parser_stack_disagreement_still_escapes_html():
    # Some older HTMLParser patch levels can keep cdata_elem after calling handle_endtag.
    parser = exporters._PdfMarkup()
    parser.handle_starttag("style", [])
    parser.set_cdata_mode("style")
    parser.handle_endtag("style")
    parser.handle_data('<link rel="prefetch" href="https://fixture.invalid">')
    assert "<link" not in "".join(parser.parts)
    assert "&lt;link" in "".join(parser.parts)
